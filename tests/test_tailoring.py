from datetime import date
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from jobagent.agents.context import render_resume
from jobagent.agents.prompts import Prompt
from jobagent.agents.tailor import Assessment, Tailor
from jobagent.agents.verifier import Verifier, decide, render_documents
from jobagent.documents.build import build_cover_letter, build_resume, month_year, plain_dashes
from jobagent.documents.docx_writer import cover_letter_docx, resume_docx
from jobagent.documents.pdf_writer import cover_letter_pdf, resume_pdf
from jobagent.llm.base import LLMError, LLMResult, Usage
from jobagent.models.candidates import Manifest
from jobagent.models.job import JOBS_HEADERS, Job, JobState
from jobagent.models.scoring import Pricing, load_scoring_config
from jobagent.models.tailoring import (
    Claim,
    ClaimCheck,
    Limits,
    TailoredBullet,
    TailoredRole,
    TailorOutput,
    VerifierOutput,
)
from jobagent.orchestrator.checks import check_tailor_output, complete_resume, numbers_in
from jobagent.orchestrator.dedupe import dedupe_key
from jobagent.orchestrator.states import InvalidTransition
from jobagent.orchestrator.tailoring import approve, retailor, run_tailoring
from jobagent.storage.csv_store import CsvStorage

from .conftest import NOW, posting

MASTER = {
    "profile": {"name": "Ada Lovelace", "location": "Toronto, ON", "phone": "555-0100", "email": "ada@example.com", "headline": "Engineering Leader"},
    "summary": "Engineering leader with 13 years of experience and 7 years in leadership.",
    "core_strengths": ["Delivery", "Architecture", "Cloud Migration"],
    "experience": [
        {"company": "Acme", "start": "2017-09", "end": "2026-09", "roles": [
            {"id": "acme-dir", "title": "Director, Engineering", "start": "2021-11", "end": "2026-09", "bullets": [
                {"id": "d-1", "text": "Led 34+ engineers across four platforms with a $15MM+ budget."},
                {"id": "d-2", "text": "Sustained 99.95% uptime on the trading platform."},
                {"id": "d-3", "text": "Led the bank's first cloud migration."}]},
            {"id": "acme-mgr", "title": "Senior Manager", "start": "2019-11", "end": "2021-11", "bullets": [
                {"id": "m-1", "text": "Led a 20-person team delivering mobile and API products."}]}]},
    ],
    "current_focus": {"title": "Applied AI", "bullets": [{"id": "ai-1", "text": "Building a RAG application with citations."}]},
    "education": [{"degree": "B.Eng.", "school": "Gujarat University"}],
}  # fmt: skip


def good_output(**kw) -> TailorOutput:
    base = {
        "summary": [Claim(text="Leader with 13 years of experience.", supports=["summary"])],
        "core_strengths": ["Delivery", "Cloud Migration"],
        "roles": [
            TailoredRole(role_id="acme-mgr", bullets=[TailoredBullet(source_id="m-1", text="Led a 20-person team on mobile and API delivery.")]),
            TailoredRole(role_id="acme-dir", bullets=[
                TailoredBullet(source_id="d-1", text="Led 34+ engineers across four platforms ($15MM+ budget)."),
                TailoredBullet(source_id="d-2", text="Sustained 99.95% uptime on the trading platform.")]),
            TailoredRole(role_id="current-focus", bullets=[TailoredBullet(source_id="ai-1", text="Building a RAG application with citations.")]),
        ],
        "cover_letter": [Claim(text="At Acme I led 34+ engineers.", supports=["d-1"]), Claim(text="I sustained 99.95% uptime.", supports=["d-2"])],
        "job_keywords_addressed": ["delivery"],
    }  # fmt: skip
    return TailorOutput(**{**base, **kw})


def verdict(*verdicts, level_ok=True, must_haves=(), feedback=()) -> VerifierOutput:
    checks = [
        ClaimCheck(claim=f"claim {i}", verdict=v, note="" if v == "traced" else "why")
        for i, v in enumerate(verdicts)
    ]
    return VerifierOutput(checks=checks, level_framing_ok=level_ok, level_framing_note="" if level_ok else "too junior",
                          missing_must_haves=list(must_haves), feedback=list(feedback))  # fmt: skip


LIMITS = Limits()


# --- cheap code checks ----------------------------------------------------------------------


def test_numbers_are_normalised():
    assert numbers_in("99.95% uptime, $15MM+ budget, 34+ engineers, 20–25%, 3x") == {
        "99.95",
        "15",
        "34",
        "20",
        "25",
        "3",
    }


def test_a_faithful_output_has_no_problems():
    assert check_tailor_output(good_output(), MASTER, LIMITS) == []


def test_unknown_ids_and_wrong_role_are_caught():
    bad = good_output(roles=[TailoredRole(role_id="acme-dir", bullets=[
        TailoredBullet(source_id="nope", text="x"), TailoredBullet(source_id="m-1", text="Led a 20-person team.")])])  # fmt: skip
    problems = " ".join(check_tailor_output(bad, MASTER, LIMITS))
    assert "unknown bullet id 'nope'" in problems and "belongs to role 'acme-mgr'" in problems
    assert "unknown role id" in " ".join(
        check_tailor_output(
            good_output(roles=[TailoredRole(role_id="zzz", bullets=[])]), MASTER, LIMITS
        )
    )


def test_an_invented_number_in_a_bullet_is_caught():
    bad = good_output(
        roles=[
            TailoredRole(
                role_id="acme-dir",
                bullets=[
                    TailoredBullet(source_id="d-1", text="Led 50+ engineers across four platforms.")
                ],
            )
        ]
    )
    problems = check_tailor_output(bad, MASTER, LIMITS)
    assert any("'50'" in p and "d-1" in p for p in problems)


def test_a_number_in_a_summary_or_letter_must_come_from_the_cited_bullets():
    bad = good_output(
        summary=[Claim(text="Leader with 15 years of experience.", supports=["summary"])],
        cover_letter=[Claim(text="I sustained 99.99% uptime.", supports=["d-2"])],
    )
    problems = " ".join(check_tailor_output(bad, MASTER, LIMITS))
    assert "'15'" in problems and "'99.99'" in problems


def test_claims_must_cite_real_ids():
    bad = good_output(
        summary=[Claim(text="Great leader.", supports=[])],
        cover_letter=[Claim(text="Hi.", supports=["ghost"])],
    )
    problems = " ".join(check_tailor_output(bad, MASTER, LIMITS))
    assert "cites no master bullet ids" in problems and "unknown ids ['ghost']" in problems


def test_core_strengths_must_be_verbatim_and_each_bullet_used_once():
    problems = " ".join(
        check_tailor_output(
            good_output(core_strengths=["Delivery", "Synergy Wizardry"]), MASTER, LIMITS
        )
    )
    assert "Synergy Wizardry" in problems
    dup = good_output(
        roles=[
            TailoredRole(
                role_id="acme-dir",
                bullets=[
                    TailoredBullet(source_id="d-1", text="a 34+"),
                    TailoredBullet(source_id="d-1", text="b 34+"),
                ],
            )
        ]
    )
    assert "used twice" in " ".join(check_tailor_output(dup, MASTER, LIMITS))


def test_length_limits_are_enforced():
    tight = Limits(
        max_bullets_per_role=1, max_total_bullets=2, max_summary_sentences=1, max_cover_paragraphs=1
    )
    problems = " ".join(
        check_tailor_output(good_output(summary=good_output().summary * 2), MASTER, tight)
    )
    for fragment in (
        "summary has 2",
        "cover letter has 2",
        "has 2 bullets (max 1)",
        "4 bullets in total (max 2)",
    ):
        assert fragment in problems


# --- Verifier decision ----------------------------------------------------------------------


def test_decision_passes_only_when_everything_is_traced_and_level_is_ok():
    assert decide(verdict("traced", "traced")).passed
    for bad in (
        verdict("traced", "inflated"),
        verdict("unsupported"),
        verdict("wrong_level"),
        verdict("traced", level_ok=False),
    ):
        d = decide(bad)
        assert not d.passed and d.problems


def test_missing_must_haves_warn_but_never_block():
    d = decide(verdict("traced", must_haves=["Kubernetes"]))
    assert d.passed and d.warnings == ["must-have not addressed: Kubernetes"]


def test_decision_feedback_names_the_failing_claims_and_the_verifier_advice():
    d = decide(verdict("inflated", feedback=["Say 'contributed', not 'led', for d-3"]))
    assert (
        any("inflated" in p for p in d.problems)
        and "Say 'contributed', not 'led', for d-3" in d.problems
    )


# --- prompts' inputs ------------------------------------------------------------------------


def test_resume_text_with_ids_lets_the_model_cite_bullets():
    text = render_resume(MASTER, include_ids=True)
    for marker in ("[summary]", "[role acme-dir]", "[role current-focus]", "[d-1]", "[ai-1]"):
        assert marker in text
    assert "[d-1]" not in render_resume(MASTER)


PRICING = Pricing(input=2.0, output=10.0, cache_read=0.2, cache_write=2.5)
PROFILES = {
    "levels": {
        "director": {"resume_emphasis": "scale and budget", "cover_letter_tone": "executive"}
    }
}
ASSESSMENT = Assessment("director", "engineering", ["manages managers"], ["no Kafka"])


def job(job_id="j1", state=JobState.APPROVED, title="Director, Platform") -> Job:
    p = posting(title=title, external_id=job_id, description="Lead engineers.")
    return Job(job_id=job_id, dedupe_key=dedupe_key(p), state=state, first_seen=NOW, posting=p)


class FakeProvider:
    """Hands out queued Tailor and Verifier outputs; a queued Exception is raised instead."""

    def __init__(self, tailor=(), verifier=()):
        self.queue = {TailorOutput: list(tailor), VerifierOutput: list(verifier)}
        self.calls: list[dict] = []

    def complete(self, **kw):
        self.calls.append(kw)
        item = self.queue[kw["schema"]].pop(0)
        if isinstance(item, Exception):
            raise item
        return LLMResult(item, Usage(1000, 500), kw["model"])


def agents(provider, cfg):
    text = render_resume(MASTER, include_ids=True)
    return (
        Tailor(provider, cfg, Prompt(1, "TAILOR"), text, PROFILES, PRICING),
        Verifier(provider, cfg, Prompt(2, "VERIFY"), text, PRICING),
    )


def test_tailor_message_carries_the_gaps_level_guidance_and_feedback(cfg):
    tailor, _ = agents(FakeProvider(), cfg)
    msg = tailor.build_message(job(), ASSESSMENT, ["fix d-1"])
    assert "<posting>" in msg and "Real level: director" in msg and "- no Kafka" in msg
    assert "Resume emphasis: scale and budget" in msg and "Cover letter tone: executive" in msg
    assert "fix d-1" in msg and "<revision_feedback>" not in tailor.build_message(
        job(), ASSESSMENT, None
    )


def test_verifier_message_shows_the_documents_with_their_citations(cfg):
    _, verifier = agents(FakeProvider(), cfg)
    msg = verifier.build_message(job(), ASSESSMENT, good_output())
    assert (
        "[d-1] Led 34+ engineers" in msg
        and "(cites: summary)" in msg
        and "<role_level>director</role_level>" in msg
    )
    assert "[C2]" in render_documents(good_output())


# --- documents ------------------------------------------------------------------------------


def test_resume_is_newest_first_with_verbatim_titles_and_dates():
    doc = build_resume(MASTER, good_output())
    assert [e.title for e in doc.experience] == [
        "Current focus: Applied AI",
        "Director, Engineering",
        "Senior Manager",
    ]
    assert doc.experience[1].dates == "Nov 2021 - Sep 2026" and doc.experience[1].company == "Acme"
    assert doc.contact == "Toronto, ON · 555-0100 · ada@example.com" and doc.education == [
        "B.Eng., Gujarat University"
    ]
    assert month_year("2019-11") == "Nov 2019"


def test_plain_dashes_replace_em_and_en_dashes():
    assert plain_dashes("Engineering — Wealth Mobile") == "Engineering - Wealth Mobile"
    assert plain_dashes("scale—and speed") == "scale - and speed"  # unspaced em dash
    assert plain_dashes("retention by 20–25%") == "retention by 20-25%"
    assert plain_dashes("Nov 2021 – Sep 2026") == "Nov 2021 - Sep 2026"
    assert plain_dashes("no dashes here, well-known") == "no dashes here, well-known"


def test_no_long_dashes_reach_any_document_even_if_the_master_or_tailor_has_them():
    master = {**MASTER, "experience": [{**MASTER["experience"][0], "roles": [
        {**MASTER["experience"][0]["roles"][0], "title": "Director — Wealth & API"}, MASTER["experience"][0]["roles"][1]]}]}  # fmt: skip
    out = good_output(summary=[Claim(text="Leader — with 13 years.", supports=["summary"])],
                      cover_letter=[Claim(text="I led 34+ engineers — at scale, 20–25% faster.", supports=["d-1"])])  # fmt: skip
    resume, letter = (
        build_resume(master, out),
        build_cover_letter(
            master, out, company="Acme", title="Director — AI", today=date(2026, 10, 7)
        ),
    )
    everything = " ".join([resume.summary, resume.headline, *resume.core_strengths, *resume.education, resume.contact]
                          + [f"{e.company} {e.title} {e.dates} {' '.join(e.bullets)}" for e in resume.experience]
                          + [letter.regarding, *letter.paragraphs])  # fmt: skip
    assert "—" not in everything and "–" not in everything
    assert "Director - Wealth & API" in everything and "20-25% faster" in everything
    pdf_text = " ".join(
        (pg.extract_text() or "") for pg in PdfReader(BytesIO(resume_pdf(resume))).pages
    )
    assert "—" not in pdf_text and "–" not in pdf_text


def test_docx_and_pdf_carry_the_same_text_and_escape_special_characters():
    out = good_output(
        summary=[Claim(text="Led R&D <teams> at 13 sites; 13 years.", supports=["summary"])]
    )
    doc = build_resume(MASTER, out)
    docx_text = "\n".join(p.text for p in Document(BytesIO(resume_docx(doc))).paragraphs)
    pdf_text = " ".join(
        (pg.extract_text() or "") for pg in PdfReader(BytesIO(resume_pdf(doc))).pages
    )
    for text in (docx_text, " ".join(pdf_text.split())):
        assert "ADA LOVELACE" in text and "Director, Engineering" in text and "R&D <teams>" in text
        assert "Led 34+ engineers across four platforms" in text


def test_cover_letter_documents():
    letter = build_cover_letter(
        MASTER, good_output(), company="Acme", title="Director, Platform", today=date(2026, 10, 7)
    )
    assert letter.date == "October 7, 2026" and letter.regarding == "Director, Platform - Acme"
    docx_text = "\n".join(p.text for p in Document(BytesIO(cover_letter_docx(letter))).paragraphs)
    pdf_text = " ".join(
        (PdfReader(BytesIO(cover_letter_pdf(letter))).pages[0].extract_text() or "").split()
    )
    assert "Dear Hiring Team," in docx_text and "At Acme I led 34+ engineers." in docx_text
    assert "At Acme I led 34+ engineers." in pdf_text


# --- run_tailoring --------------------------------------------------------------------------


class FakeStore:
    """An in-memory Drive: folders, files, downloads and trash."""

    def __init__(self):
        self.folders: dict[str, str] = {}  # folder id -> name
        self.files: list[tuple[str, str]] = []  # (folder id, file name)
        self.data: dict[str, bytes] = {}  # file name -> latest bytes
        self.by_id: dict[str, bytes] = {}  # file id -> bytes
        self.trashed: list[str] = []

    def _folder(self, name):
        fid = f"folder{len(self.folders) + 1}"
        self.folders[fid] = name
        return fid, f"https://drive.example/{fid}"

    def create_job_folder(self, name):
        return self._folder(name)

    def create_subfolder(self, name, parent_id):
        assert parent_id in self.folders, "subfolder needs an existing parent"
        return self._folder(name)

    def upload_bytes(self, name, data, mime_type, parent_id=None):
        assert data, f"{name} is empty"
        self.files.append((parent_id, name))
        self.data[name] = data
        file_id = f"file{len(self.files)}"
        self.by_id[file_id] = data
        return file_id

    def download_bytes(self, file_id):
        return self.by_id[file_id]

    def trash(self, file_id):
        self.trashed.append(file_id)

    def names_in(self, folder_name):
        ids = [i for i, n in self.folders.items() if n == folder_name]
        return sorted(n for f, n in self.files if f in ids)


def storage_with(tmp_path, *jobs: Job) -> CsvStorage:
    s = CsvStorage(str(tmp_path))
    s.append_records(
        "Jobs",
        JOBS_HEADERS,
        [
            j.to_record()
            | {"real_level": "director", "function": "engineering", "gaps": "no Kafka"}
            for j in jobs
        ],
    )
    return s


def run(storage, provider, cfg, store=None, **kw):
    tailor, verifier = agents(provider, cfg)
    store = store or FakeStore()
    return run_tailoring(
        storage, tailor, verifier, store, MASTER, cfg, today=date(2026, 10, 7), **kw
    ), store


def only_row(storage):
    [row] = storage.read_records("Jobs")
    return row


def test_happy_path_moves_to_ready_and_saves_resume_and_report(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    summary, store = run(
        storage,
        FakeProvider([good_output()], [verdict("traced", "traced", must_haves=["Kafka"])]),
        cfg,
    )

    row = only_row(storage)
    assert (
        row["state"] == "ready"
        and row["tailor_status"] == "ready"
        and row["tailor_attempts"] == "1"
    )
    assert row["drive_url"] == "https://drive.example/folder2" and "Kafka" in row["tailor_notes"]
    assert store.names_in("candidate-1") == [
        "resume.docx",
        "resume.pdf",
        "tailored.json",
        "verifier_report.md",
    ]
    assert (
        row["final_candidate"] == "1"
        and Manifest.from_cell(row["tailor_candidates"]).candidates[0].n == 1
    )
    assert summary.outcomes[0].status == "ready" and summary.cost_usd == pytest.approx(
        2 * (1000 * 2 + 500 * 10) / 1e6
    )
    assert list(store.folders.values()) == ["Acme - Director, Platform - j1", "candidate-1"]


def test_code_check_failure_goes_back_to_the_tailor_without_calling_the_verifier(tmp_path, cfg):
    invented = good_output(
        roles=[
            TailoredRole(
                role_id="acme-dir",
                bullets=[TailoredBullet(source_id="d-1", text="Led 50+ engineers.")],
            )
        ]
    )
    provider = FakeProvider([invented, good_output()], [verdict("traced")])
    storage = storage_with(tmp_path, job())
    run(storage, provider, cfg)

    assert only_row(storage)["state"] == "ready" and only_row(storage)["tailor_attempts"] == "2"
    assert [c["schema"] for c in provider.calls] == [TailorOutput, TailorOutput, VerifierOutput]
    assert "'50'" in provider.calls[1]["prompt"]  # the exact problem was sent back


def test_verifier_failure_sends_the_tailor_back_once_then_passes(tmp_path, cfg):
    provider = FakeProvider(
        [good_output(), good_output()],
        [verdict("inflated", feedback=["tone it down"]), verdict("traced")],
    )
    storage = storage_with(tmp_path, job())
    run(storage, provider, cfg)

    assert only_row(storage)["state"] == "ready" and only_row(storage)["tailor_attempts"] == "2"
    assert "tone it down" in provider.calls[2]["prompt"]


def test_two_verifier_failures_block_the_job_and_it_stays_approved(tmp_path, cfg):
    provider = FakeProvider(
        [good_output(), good_output()], [verdict("unsupported"), verdict("wrong_level")]
    )
    storage = storage_with(tmp_path, job())
    summary, store = run(storage, provider, cfg)

    row = only_row(storage)
    assert (
        row["state"] == "approved"
        and row["tailor_status"] == "blocked"
        and row["tailor_attempts"] == "2"
    )
    assert row["tailor_notes"] and summary.outcomes[0].status == "blocked"
    assert store.files == [
        ("folder1", "blocked_report.md")
    ]  # in the job folder; no candidate, no resume
    assert (
        Manifest.from_cell(row["tailor_candidates"]).candidates == []
        and row["drive_url"] == "https://drive.example/folder1"
    )


def test_blocked_jobs_are_skipped_unless_retried(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    drive = FakeStore()
    run(
        storage,
        FakeProvider(
            [good_output(), good_output()], [verdict("unsupported"), verdict("unsupported")]
        ),
        cfg,
        drive,
    )

    skipped, _ = run(storage, FakeProvider(), cfg, drive)
    assert skipped.outcomes == []
    again, _ = run(
        storage, FakeProvider([good_output()], [verdict("traced")]), cfg, drive, retry_blocked=True
    )
    assert again.outcomes[0].status == "ready" and only_row(storage)["state"] == "ready"


def test_a_model_error_leaves_the_job_approved_for_the_next_run(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    summary, store = run(storage, FakeProvider([LLMError("overloaded")]), cfg)
    assert summary.outcomes[0].status == "error" and store.files == []
    row = only_row(storage)
    assert (
        row["state"] == "approved"
        and row["tailor_status"] == "error"
        and "overloaded" in row["tailor_notes"]
    )
    retry, _ = run(
        storage, FakeProvider([good_output()], [verdict("traced")]), cfg
    )  # error is not 'blocked'
    assert retry.outcomes[0].status == "ready"


def test_only_approved_jobs_are_tailored(tmp_path, cfg):
    storage = storage_with(
        tmp_path,
        job("a", JobState.SHORTLISTED),
        job("b", JobState.APPROVED),
        job("c", JobState.READY),
    )
    provider = FakeProvider([good_output()], [verdict("traced")])
    summary, _ = run(storage, provider, cfg)
    assert [o.job_id for o in summary.outcomes] == ["b"]


def test_spend_cap_stops_before_the_next_job(tmp_path, cfg):
    capped = cfg.model_copy(update={"max_run_cost_usd": 0.001})
    storage = storage_with(tmp_path, job("a"), job("b"))
    provider = FakeProvider([good_output(), good_output()], [verdict("traced"), verdict("traced")])
    summary, _ = run(storage, provider, capped)
    assert [o.job_id for o in summary.outcomes] == ["a"] and "spend cap" in summary.stopped


# --- the human gate -------------------------------------------------------------------------


def test_only_shortlisted_jobs_can_be_approved_and_only_via_the_human_path(tmp_path):
    storage = storage_with(
        tmp_path,
        job("s", JobState.SHORTLISTED),
        job("f", JobState.FILTERED),
        job("r", JobState.READY),
    )
    results = approve(storage, ["s", "f", "r", "ghost"])
    assert results["s"].startswith("approved") and results["ghost"] == "not found"
    assert results["f"].startswith("not approved") and results["r"].startswith("not approved")
    states = {r["job_id"]: r["state"] for r in storage.read_records("Jobs")}
    assert states == {"s": "approved", "f": "filtered", "r": "ready"}
    with pytest.raises(InvalidTransition):  # the orchestrator itself can never approve
        from jobagent.orchestrator.states import transition

        transition(job("x", JobState.SHORTLISTED), JobState.APPROVED)


def test_pricing_for_the_configured_models_exists(cfg):
    prices = load_scoring_config().pricing_usd_per_mtok
    assert cfg.tailor.model in prices and cfg.verifier.model in prices


# --- cover letter switch ------------------------------------------------------------------


def test_cover_letter_is_off_by_default_in_the_config(cfg):
    assert cfg.cover_letter is False


def test_tailor_is_told_when_no_cover_letter_is_wanted_and_asked_for_one_when_enabled(cfg):
    off, _ = agents(FakeProvider(), cfg)
    on, _ = agents(FakeProvider(), cfg.model_copy(update={"cover_letter": True}))
    assert "Return an empty cover_letter list" in off.build_message(job(), ASSESSMENT, None)
    assert "<cover_letter>" not in on.build_message(job(), ASSESSMENT, None)


def test_a_letter_the_model_writes_anyway_is_dropped_before_checking_and_saving(tmp_path, cfg):
    bad_letter = good_output(cover_letter=[Claim(text="I led 999 engineers.", supports=["d-1"])])
    provider = FakeProvider([bad_letter], [verdict("traced")])
    storage = storage_with(tmp_path, job())
    summary, store = run(storage, provider, cfg)

    assert summary.outcomes[0].status == "ready"  # the bad letter would have failed the code checks
    assert "cover_letter.pdf" not in [n for _, n in store.files]
    assert "[C1]" not in provider.calls[1]["prompt"]  # the Verifier never saw a letter


def test_enabling_the_switch_saves_the_letter_files(tmp_path, cfg):
    on = cfg.model_copy(update={"cover_letter": True})
    storage = storage_with(tmp_path, job())
    summary, store = run(storage, FakeProvider([good_output()], [verdict("traced")]), on)
    assert summary.outcomes[0].status == "ready"
    assert store.names_in("candidate-1") == [
        "cover_letter.docx",
        "cover_letter.pdf",
        "resume.docx",
        "resume.pdf",
        "tailored.json",
        "verifier_report.md",
    ]


# --- keep all roles -----------------------------------------------------------------------


def only_director(**kw) -> TailorOutput:
    """The Tailor tailors the Director role and leaves the rest out, as it tends to."""
    return good_output(
        roles=[
            TailoredRole(
                role_id="acme-dir",
                bullets=[
                    TailoredBullet(
                        source_id="d-1",
                        text="Led 34+ engineers across four platforms ($15MM+ budget).",
                    )
                ],
            )
        ],
        **kw,
    )


def test_all_roles_adds_missing_roles_verbatim_and_leaves_the_rest_alone():
    tailored = only_director()
    done, added, extra = complete_resume(tailored, MASTER, "all_roles")

    assert (added, extra) == (["acme-mgr", "current-focus"], 0)
    by_id = {r.role_id: r for r in done.roles}
    assert [b.text for b in by_id["acme-mgr"].bullets] == [
        "Led a 20-person team delivering mobile and API products."
    ]
    assert [b.source_id for b in by_id["current-focus"].bullets] == ["ai-1"]
    assert by_id["acme-dir"] == tailored.roles[0]  # all_roles never pads a role the Tailor included
    assert complete_resume(done, MASTER, "all_roles")[1:] == (
        [],
        0,
    )  # nothing left to add the second time


def test_all_bullets_appends_missing_bullets_after_the_tailors_own_in_each_role():
    tailored = only_director()  # d-1 reworded; d-2 and d-3 left out
    done, added, extra = complete_resume(tailored, MASTER, "all_bullets")

    director = next(r for r in done.roles if r.role_id == "acme-dir")
    assert [b.source_id for b in director.bullets] == [
        "d-1",
        "d-2",
        "d-3",
    ]  # tailored first, rest in master order
    assert (
        director.bullets[0].text == "Led 34+ engineers across four platforms ($15MM+ budget)."
    )  # rewording kept
    assert (
        director.bullets[1].text == "Sustained 99.95% uptime on the trading platform."
    )  # verbatim
    assert (added, extra) == (["acme-mgr", "current-focus"], 2)
    every = {b.source_id for r in done.roles for b in r.bullets}
    assert every == {"d-1", "d-2", "d-3", "m-1", "ai-1"}  # the whole master


def test_tailor_choice_changes_nothing():
    tailored = only_director()
    assert complete_resume(tailored, MASTER, "tailor_choice") == (tailored, [], 0)


def test_bullet_limits_apply_only_when_the_tailor_decides_the_content():
    tight = Limits(max_total_bullets=1)
    assert any("bullets in total" in p for p in check_tailor_output(good_output(), MASTER, tight))
    assert check_tailor_output(good_output(), MASTER, tight, enforce_bullet_limits=False) == []


def read_docx(data: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)


def test_the_saved_resume_keeps_every_role_with_default_text(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    summary, store = run(
        storage,
        FakeProvider([only_director()], [verdict("traced")]),
        cfg.model_copy(update={"keep": "all_roles"}),
    )

    text = read_docx(store.data["resume.docx"])
    assert (
        "Senior Manager" in text
        and "Led a 20-person team delivering mobile and API products." in text
    )  # default text
    assert "Building a RAG application with citations." in text
    assert "added unchanged from the master: 2 roles" in store.data["verifier_report.md"].decode()
    assert summary.outcomes[0].status == "ready"


def test_tailor_choice_keeps_only_what_the_tailor_selected(tmp_path, cfg):
    off = cfg.model_copy(update={"keep": "tailor_choice"})
    storage = storage_with(tmp_path, job())
    _, store = run(storage, FakeProvider([only_director()], [verdict("traced")]), off)
    text = read_docx(store.data["resume.docx"])
    assert "Senior Manager" not in text and "Director, Engineering" in text


def test_verifier_does_not_see_the_roles_added_after_it_passed(tmp_path, cfg):
    provider = FakeProvider([only_director()], [verdict("traced")])
    run(storage_with(tmp_path, job()), provider, cfg)
    assert "20-person team" not in provider.calls[1]["prompt"]


def test_default_keeps_every_bullet_and_never_sends_the_added_ones_to_the_verifier(tmp_path, cfg):
    assert cfg.keep == "all_bullets"
    provider = FakeProvider([only_director()], [verdict("traced")])
    summary, store = run(storage_with(tmp_path, job()), provider, cfg)

    text = read_docx(store.data["resume.docx"])
    for fact in (
        "Led the bank's first cloud migration.",
        "Sustained 99.95% uptime",
        "20-person team",
        "RAG application",
    ):
        assert fact in text
    assert "first cloud migration" not in provider.calls[1]["prompt"]  # copied after verification
    assert "2 bullets in other roles" in store.data["verifier_report.md"].decode()
    assert summary.outcomes[0].status == "ready"


# --- retailor (human gate for a second try) ------------------------------------------------


def test_ready_job_goes_back_to_approved_and_its_tailoring_columns_are_cleared(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    drive = FakeStore()
    run(storage, FakeProvider([good_output()], [verdict("traced")]), cfg, drive)
    assert only_row(storage)["state"] == "ready"

    result = retailor(storage, ["j1"])

    row = only_row(storage)
    assert result["j1"].startswith("approved again")
    assert row["state"] == "approved" and row["tailor_status"] == ""
    assert (
        row["drive_url"] and len(Manifest.from_cell(row["tailor_candidates"]).candidates) == 1
    )  # kept
    again, _ = run(
        storage, FakeProvider([good_output()], [verdict("traced")]), cfg, drive
    )  # picked up without --retry
    assert again.outcomes[0].status == "ready"


def test_blocked_job_is_reset_so_tailor_picks_it_up_without_retry_flag(tmp_path, cfg):
    storage = storage_with(tmp_path, job())
    run(storage, FakeProvider([good_output(), good_output()], [verdict("unsupported")] * 2), cfg)
    assert only_row(storage)["tailor_status"] == "blocked"

    assert retailor(storage, ["j1"])["j1"].startswith("reset")
    assert only_row(storage)["state"] == "approved" and only_row(storage)["tailor_status"] == ""


def test_retailor_refuses_jobs_that_were_never_tailored_or_were_submitted(tmp_path):
    storage = storage_with(
        tmp_path,
        job("s", JobState.SHORTLISTED),
        job("x", JobState.SUBMITTED),
        job("a", JobState.APPROVED),
    )
    results = retailor(storage, ["s", "x", "a", "ghost"])
    assert results["ghost"] == "not found"
    for refused in ("s", "x", "a"):
        assert results[refused].startswith("not changed")
    states = {r["job_id"]: r["state"] for r in storage.read_records("Jobs")}
    assert states == {"s": "shortlisted", "x": "submitted", "a": "approved"}


def test_only_a_human_can_move_ready_back_to_approved():
    from jobagent.orchestrator.states import transition

    with pytest.raises(InvalidTransition):
        transition(job("r", JobState.READY), JobState.APPROVED)  # the orchestrator may not
    assert (
        transition(job("r", JobState.READY), JobState.APPROVED, actor="human").state
        == JobState.APPROVED
    )
    with pytest.raises(InvalidTransition):  # a submitted job never returns to approved
        transition(job("x", JobState.SUBMITTED), JobState.APPROVED, actor="human")
