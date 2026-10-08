from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from jobagent.documents.model import CoverLetterDoc, ResumeDoc

FONT = "Calibri"


def _base() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(2)
    return doc


def _rule(paragraph) -> None:
    """A thin line under a heading: a paragraph border, so no table or shape is needed."""
    pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(k), v)
    border.append(bottom)
    pr.append(border)


def _centered(doc, text: str, *, size: float, bold: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size, run.bold = Pt(size), bold
    return p


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(text.upper())
    run.bold, run.font.size = True, Pt(11)
    _rule(p)


def _save(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def resume_docx(d: ResumeDoc) -> bytes:
    doc = _base()
    _centered(doc, d.name.upper(), size=18, bold=True)
    _centered(doc, d.contact, size=9.5)
    if d.headline:
        _centered(doc, d.headline, size=10, bold=True)
    _heading(doc, "Summary")
    doc.add_paragraph(d.summary)
    _heading(doc, "Core strengths")
    doc.add_paragraph(" · ".join(d.core_strengths))
    _heading(doc, "Professional experience")
    for e in d.experience:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = True
        head = f"{e.company}  |  " if e.company else ""
        p.add_run(head).bold = True
        r = p.add_run(e.title)
        r.bold = r.italic = True
        if e.dates:
            p.add_run(f"  |  {e.dates}")
        for b in e.bullets:
            doc.add_paragraph(b, style="List Bullet")
    if d.education:
        _heading(doc, "Education")
        for line in d.education:
            doc.add_paragraph(line)
    return _save(doc)


def cover_letter_docx(d: CoverLetterDoc) -> bytes:
    doc = _base()
    _centered(doc, d.name.upper(), size=16, bold=True)
    _centered(doc, d.contact, size=9.5)
    doc.add_paragraph()
    doc.add_paragraph(d.date)
    p = doc.add_paragraph()
    p.add_run("Re: ").bold = True
    p.add_run(d.regarding)
    doc.add_paragraph()
    doc.add_paragraph(d.greeting)
    for para in d.paragraphs:
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(8)
    doc.add_paragraph(d.closing)
    doc.add_paragraph(d.name)
    return _save(doc)
